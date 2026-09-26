"""
Unit tests for CSV and Excel export functionality (Prompt 5).
Validates file generation, headers, Unicode handling, missing value representation,
and FastAPI export endpoints.
"""
import io
import csv
from unittest.mock import patch
import openpyxl
import pytest
from starlette.testclient import TestClient

from app.main import app
from app.export.exporter import generate_csv, generate_excel


SAMPLE_GROUPS = [
    {
        "id": "11111111-1111-1111-1111-111111111111",
        "name": "NYC Tech & AI Founders 🚀",
        "url": "https://www.facebook.com/groups/nyctech",
        "member_count": 12500,
        "privacy": "Public",
        "niche": "Tech",
        "country": "USA",
        "matched_keywords": ["tech", "ai", "nyc"],
        "activity_status": "Active",
        "external_link_status": "Allowed",
        "discovered_at": "2026-09-20T10:00:00Z",
        "analysis": {
            "activity_status": "Active",
            "activity_score": 88,
            "external_link_status": "Allowed",
            "created_at": "2026-09-20T12:00:00Z",
        },
    },
    {
        "id": "22222222-2222-2222-2222-222222222222",
        "name": "London Singles Dating",
        "url": "https://www.facebook.com/groups/londonsingles",
        "member_count": None,
        "privacy": "Private",
        "niche": "Dating",
        "country": "UK",
        "matched_keywords": ["dating", "singles"],
        "activity_status": None,
        "external_link_status": None,
        "discovered_at": "2026-09-21T15:30:00Z",
        "analysis": None,
    },
]


def test_generate_csv_headers_and_data():
    """CSV export includes expected headers and maps record fields accurately."""
    csv_text = generate_csv(SAMPLE_GROUPS)
    assert isinstance(csv_text, str)

    reader = list(csv.reader(io.StringIO(csv_text)))
    assert len(reader) == 3  # Header + 2 data rows
    headers = reader[0]
    expected_headers = [
        "Group Name",
        "Facebook URL",
        "Member Count",
        "Privacy",
        "Niche",
        "Country",
        "Matched Keywords",
        "Activity Status",
        "Activity Score",
        "External Link Status",
        "Discovery Date",
        "Analysis Date",
    ]
    # Check that required columns are in headers
    for h in expected_headers:
        assert h in headers

    row1 = reader[1]
    assert row1[0] == "NYC Tech & AI Founders 🚀"
    assert row1[1] == "https://www.facebook.com/groups/nyctech"
    assert row1[2] == "12500"
    assert row1[3] == "Public"
    assert row1[4] == "Tech"
    assert row1[5] == "USA"
    assert "tech" in row1[6]
    assert row1[7] == "Active"
    assert row1[8] == "88"
    assert row1[9] == "Allowed"

    # Second row with missing values -> display Unknown / N/A, not blank or invented
    row2 = reader[2]
    assert row2[0] == "London Singles Dating"
    assert row2[2] == "Unknown"
    assert row2[7] == "Unknown"
    assert row2[8] == "N/A"
    assert row2[9] == "Unknown"
    assert row2[11] == "N/A"


def test_generate_csv_empty_dataset():
    """Generating CSV for empty dataset returns a valid CSV with only headers."""
    csv_text = generate_csv([])
    assert isinstance(csv_text, str)
    reader = list(csv.reader(io.StringIO(csv_text)))
    assert len(reader) == 1
    assert "Group Name" in reader[0]


def test_generate_excel_valid_xlsx():
    """Excel export generates a valid .xlsx spreadsheet readable by openpyxl."""
    excel_bytes = generate_excel(SAMPLE_GROUPS)
    assert isinstance(excel_bytes, bytes)
    # XLSX files start with the PK zip header
    assert excel_bytes[:2] == b"PK"

    # Load with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet = wb.active
    assert sheet.title == "Facebook Groups"

    # Verify header
    headers = [cell.value for cell in sheet[1]]
    assert "Group Name" in headers
    assert "Facebook URL" in headers
    assert "Activity Score" in headers

    # Verify rows
    assert sheet.max_row == 3
    row1_cells = [cell.value for cell in sheet[2]]
    assert row1_cells[0] == "NYC Tech & AI Founders 🚀"
    assert str(row1_cells[2]) == "12500"

    row2_cells = [cell.value for cell in sheet[3]]
    assert row2_cells[0] == "London Singles Dating"
    assert row2_cells[2] == "Unknown"
    assert row2_cells[8] == "N/A"


def test_generate_excel_empty():
    """Generating Excel for empty list produces a sheet with headers only."""
    excel_bytes = generate_excel([])
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet = wb.active
    assert sheet.max_row == 1


# ==============================================================================
# API Endpoint Tests
# ==============================================================================

def test_api_export_csv_endpoint():
    """GET /api/export/csv returns text/csv with attachment header."""
    client = TestClient(app)

    with patch("app.main.query_groups_filtered") as mock_query:
        mock_query.return_value = {
            "total": 1,
            "page": 1,
            "page_size": 1,
            "total_pages": 1,
            "groups": [SAMPLE_GROUPS[0]],
        }

        res = client.get("/api/export/csv?niche=Tech")
        assert res.status_code == 200
        assert "text/csv" in res.headers["content-type"]
        assert "attachment" in res.headers["content-disposition"]
        assert "facebook_groups_" in res.headers["content-disposition"]
        assert res.headers["content-disposition"].endswith('.csv"')
        assert "NYC Tech & AI Founders" in res.text


def test_api_export_excel_endpoint():
    """GET /api/export/excel returns application/vnd.openxmlformats-officedocument.spreadsheetml.sheet."""
    client = TestClient(app)

    with patch("app.main.query_groups_filtered") as mock_query:
        mock_query.return_value = {
            "total": 1,
            "page": 1,
            "page_size": 1,
            "total_pages": 1,
            "groups": [SAMPLE_GROUPS[0]],
        }

        res = client.get("/api/export/excel?country=USA")
        assert res.status_code == 200
        assert "spreadsheetml.sheet" in res.headers["content-type"]
        assert "attachment" in res.headers["content-disposition"]
        assert "facebook_groups_" in res.headers["content-disposition"]
        assert res.headers["content-disposition"].endswith('.xlsx"')
        assert res.content[:2] == b"PK"
