"""
CSV and Excel export utilities for Facebook Group Research Agent.
Formats and serializes filtered group research data with attached intelligence.
"""
import csv
import io
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CSV_COLUMNS = [
    ("name", "Group Name"),
    ("facebook_url", "Facebook URL"),
    ("member_count", "Member Count"),
    ("privacy", "Privacy"),
    ("niche", "Niche"),
    ("country", "Country"),
    ("matched_keywords", "Matched Keywords"),
    ("activity_status", "Activity Status"),
    ("activity_score", "Activity Score"),
    ("external_link_status", "External Link Status"),
    ("discovered_at", "Discovery Date"),
    ("analyzed_at", "Analysis Date"),
    ("rules_summary", "Rules Summary"),
    ("activity_summary", "Activity Summary"),
    ("overall_summary", "Overall Summary"),
    ("external_link_evidence", "External Link Evidence"),
]


def _format_cell_value(val: Any) -> str:
    """Formats cell values cleanly without inventing data."""
    if val is None:
        return "Unknown"
    if isinstance(val, list):
        if not val:
            return ""
        return ", ".join(str(item) for item in val)
    if isinstance(val, (int, float)):
        return str(val)
    val_str = str(val).strip()
    return val_str if val_str else "Unknown"


def _flatten_group_row(group: Dict[str, Any]) -> Dict[str, str]:
    """Flattens group and nested analysis record into a standardized dictionary."""
    # Check if analysis is embedded
    analysis: Optional[Dict[str, Any]] = None
    analyses_list = group.get("analyses")
    if isinstance(analyses_list, list) and len(analyses_list) > 0:
        analysis = analyses_list[0]
    elif isinstance(group.get("analysis"), dict):
        analysis = group.get("analysis")

    # Member count formatting
    member_count_val = group.get("member_count")
    if member_count_val is None:
        member_display = group.get("member_count_text") or "Unknown"
    else:
        member_display = str(member_count_val)

    # Keywords formatting
    kws = group.get("matched_keywords")
    if not kws and group.get("keyword"):
        kws = [group.get("keyword")]

    # Activity status
    act_status = None
    if analysis and analysis.get("activity_status"):
        act_status = analysis.get("activity_status")
    elif group.get("activity_status"):
        act_status = group.get("activity_status")

    # External link status
    ext_status = None
    if analysis and analysis.get("external_link_status"):
        ext_status = analysis.get("external_link_status")
    elif group.get("external_link_status"):
        ext_status = group.get("external_link_status")

    # Analysis date
    analyzed_date = None
    if analysis:
        analyzed_date = analysis.get("analyzed_at") or analysis.get("created_at")

    row = {
        "name": group.get("name") or "Unknown",
        "facebook_url": group.get("facebook_url") or group.get("url") or "Unknown",
        "member_count": member_display,
        "privacy": group.get("privacy") or "Unknown",
        "niche": group.get("niche") or "Unknown",
        "country": group.get("country") or "Unknown",
        "matched_keywords": _format_cell_value(kws) if kws else "Unknown",
        "discovered_at": group.get("discovered_at") or "Unknown",
        "activity_status": act_status or "Unknown",
        "activity_score": str(analysis.get("activity_score")) if (analysis and analysis.get("activity_score") is not None) else "N/A",
        "external_link_status": ext_status or "Unknown",
        "analyzed_at": analyzed_date or "N/A",
        "rules_summary": (analysis.get("rules_summary") if analysis else "") or "",
        "activity_summary": (analysis.get("activity_summary") if analysis else "") or "",
        "overall_summary": (analysis.get("overall_summary") if analysis else (analysis.get("summary") if analysis else "")) or "",
        "external_link_evidence": (analysis.get("external_link_evidence") if analysis else "") or "",
    }
    return row


def generate_csv(groups: List[Dict[str, Any]]) -> str:
    """
    Generates a RFC-4180 compliant CSV string from the list of group records.
    Handles Unicode encoding and empty data gracefully.
    """
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

    # Write header row
    headers = [col_title for _, col_title in CSV_COLUMNS]
    writer.writerow(headers)

    # Write data rows
    for group in groups:
        row_dict = _flatten_group_row(group)
        row_values = [row_dict.get(col_key, "") for col_key, _ in CSV_COLUMNS]
        writer.writerow(row_values)

    return output.getvalue()


def generate_excel(groups: List[Dict[str, Any]]) -> bytes:
    """
    Generates a styled Excel workbook (.xlsx) binary from the list of group records.
    Uses openpyxl with formatted headers, auto-adjusted column widths, and borders.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facebook Groups"

    # Header styling
    header_fill = PatternFill(start_color="1877F2", end_color="1877F2", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Write header row
    headers = [col_title for _, col_title in CSV_COLUMNS]
    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Write data rows
    for row_idx, group in enumerate(groups, start=2):
        row_dict = _flatten_group_row(group)
        row_values = [row_dict.get(col_key, "") for col_key, _ in CSV_COLUMNS]
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 20

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = cell_alignment

    # Auto-adjust column widths
    for col_idx, (col_key, col_title) in enumerate(CSV_COLUMNS, start=1):
        col_letter = get_column_letter(col_idx)
        # Calculate max width
        max_len = len(col_title)
        for row in range(2, ws.max_row + 1):
            val = str(ws.cell(row=row, column=col_idx).value or "")
            if len(val) > max_len:
                max_len = min(len(val), 50)  # cap width
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
